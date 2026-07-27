from io import BytesIO
from pathlib import Path
import sys
import unittest

from docx import Document as WordDocument

APP_DIR = Path(__file__).resolve().parents[1] / "app"
PROJECT_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(APP_DIR))

from models.review import ReviewStatus
from services.task_service import create_review_task
from tools.registry import tool_registry


class DocumentPipelineTest(unittest.TestCase):
    def test_docx_upload_parses_paragraphs_tables_and_clauses(self):
        task = create_review_task(
            file_name="sample.docx",
            content=build_docx_bytes(),
            review_position_value="甲方",
        )

        self.assertEqual(task.status, ReviewStatus.EVIDENCE_VERIFIED)
        payload = task.to_dict()
        self.assertEqual(payload["file_type"], "docx")
        self.assertEqual(len(payload["document"]["paragraphs"]), 4)
        self.assertEqual(len(payload["document"]["tables"]), 1)
        self.assertTrue(payload["document"]["content_hash"])
        self.assertEqual(payload["clauses"][0]["clause_id"], "CL-001")
        self.assertEqual(payload["clauses"][0]["clause_type"], "定义")
        self.assertIn("义务", payload["clauses"][1]["title"])
        self.assertEqual(payload["clauses"][1]["clause_type"], "保密义务")
        self.assertIn("接收方", payload["clauses"][1]["key_fields"]["obligation_subject"])
        self.assertIn("3年", payload["clauses"][1]["key_fields"]["confidentiality_period"])
        self.assertEqual(
            [log["tool_name"] for log in payload["logs"][:3]],
            ["parse_document", "classify_contract_type", "extract_clauses"],
        )
        self.assertEqual(payload["matched_rules"][0]["matched_rules"][0]["rule_id"], "NDA-R001")
        self.assertEqual(payload["review_contexts"][0]["matched_rule"]["rule_id"], "NDA-R001")
        self.assertFalse(payload["review_contexts"][0]["formal_risk_generated"])
        self.assertEqual(payload["risk_findings"][0]["risk_type"], "保密信息范围过宽")
        self.assertIn("商业信息", payload["risk_findings"][0]["evidence_text"])

    def test_pdf_upload_extracts_readable_text(self):
        task = create_review_task(
            file_name="sample.pdf",
            content=build_pdf_bytes(),
            review_position_value="乙方",
        )

        payload = task.to_dict()
        self.assertEqual(task.status, ReviewStatus.HUMAN_REVIEW_PENDING)
        self.assertEqual(payload["file_type"], "pdf")
        self.assertIn(
            "Confidential Information",
            "\n".join(paragraph["text"] for paragraph in payload["document"]["paragraphs"]),
        )
        self.assertEqual(len(payload["clauses"]), 2)
        self.assertEqual(payload["clauses"][1]["clause_type"], "期限")
        self.assertEqual(payload["matched_rules"][1]["matched_rules"][0]["rule_id"], "NDA-R003")
        self.assertTrue(payload["review_contexts"])
        self.assertTrue(payload["analysis_results"])
        self.assertEqual(payload["document"]["paragraphs"][0]["source_location"]["page_number"], 1)
        self.assertEqual(payload["document"]["paragraphs"][-1]["source_location"]["page_number"], 2)
        self.assertEqual(len(payload["document"]["paragraphs"][0]["source_location"]["bbox"]), 4)
        self.assertEqual(payload["clauses"][1]["source_location"]["start_page"], 2)
        self.assertEqual(
            payload["clauses"][1]["source_location"]["pdf_blocks"][0]["block_id"],
            payload["document"]["paragraphs"][2]["block_id"],
        )
        risk_clause = next(
            clause
            for clause in payload["clauses"]
            if clause["clause_id"] == payload["risk_findings"][0]["clause_id"]
        )
        self.assertIn(payload["risk_findings"][0]["evidence_text"], risk_clause["text"])
        self.assertTrue(risk_clause["source_location"]["pdf_blocks"])
        evidence_location = payload["risk_findings"][0]["evidence_location"]
        self.assertEqual(evidence_location["start_page"], 2)
        self.assertEqual(evidence_location["end_page"], 2)
        self.assertEqual(
            [block["block_id"] for block in evidence_location["pdf_blocks"]],
            ["B0003", "B0004"],
        )
        self.assertEqual(evidence_location["pdf_blocks"][0]["evidence_start"], 3)
        self.assertEqual(len(evidence_location["pdf_blocks"][0]["bbox"]), 4)
        self.assertEqual(payload["evidence_results"][0]["source_location"], evidence_location)

    def test_scanned_pdf_upload_uses_ocr_before_review(self):
        fixture = (
            PROJECT_ROOT
            / "backend"
            / "tests"
            / "fixtures"
            / "pdf"
            / "scanned_image.pdf"
        )

        task = create_review_task(
            file_name=fixture.name,
            content=fixture.read_bytes(),
            review_position_value="甲方",
        )

        payload = task.to_dict()
        self.assertNotEqual(task.status, ReviewStatus.PARSE_FAILED)
        self.assertEqual(
            payload["document"]["blocks"][0]["text"],
            "SCANNED NDA IMAGE - NO PDF TEXT LAYER",
        )
        self.assertEqual(
            payload["document"]["blocks"][0]["source_location"]["page_number"],
            1,
        )

    def test_pdf_parse_failures_stop_before_risk_analysis(self):
        cases = (
            ("empty_text.pdf", "OCR"),
            ("encrypted.pdf", "加密"),
            ("complex_font.pdf", "复杂字体"),
            ("damaged.pdf", "损坏"),
        )
        fixtures = PROJECT_ROOT / "backend" / "tests" / "fixtures" / "pdf"

        for file_name, message_fragment in cases:
            with self.subTest(file_name=file_name):
                task = create_review_task(
                    file_name=file_name,
                    content=(fixtures / file_name).read_bytes(),
                    review_position_value="甲方",
                )

                payload = task.to_dict()
                self.assertEqual(task.status, ReviewStatus.PARSE_FAILED)
                self.assertIn(message_fragment, payload["message"])
                self.assertIsNone(payload["document"])
                self.assertIsNone(payload["clauses"])
                self.assertIsNone(payload["risk_findings"])
                self.assertEqual(payload["logs"][0]["tool_name"], "parse_document")
                self.assertEqual(payload["logs"][0]["status"], "failed")

    def test_confidentiality_obligation_is_not_misclassified_as_definition(self):
        task = create_review_task(
            file_name="sample.docx",
            content=build_confidentiality_docx_bytes(),
            review_position_value="甲方",
        )

        payload = task.to_dict()
        self.assertEqual(payload["clauses"][0]["clause_type"], "保密义务")

    def test_unreadable_docx_returns_parse_failed(self):
        task = create_review_task(
            file_name="broken.docx",
            content=b"not a docx",
            review_position_value="甲方",
        )

        payload = task.to_dict()
        self.assertEqual(task.status, ReviewStatus.PARSE_FAILED)
        self.assertEqual(payload["logs"][0]["tool_name"], "parse_document")
        self.assertEqual(payload["logs"][0]["status"], "failed")
        self.assertIsNone(payload["document"])
        self.assertIsNone(payload["clauses"])

    def test_rejects_unsupported_file_type(self):
        with self.assertRaises(ValueError):
            create_review_task(
                file_name="sample.exe",
                content=b"hello",
                review_position_value="甲方",
            )

    def test_task_two_tools_are_registered(self):
        self.assertIn("parse_document", tool_registry)
        self.assertIn("extract_clauses", tool_registry)
        self.assertIn("extract_key_fields", tool_registry)
        self.assertIn("retrieve_playbook_rules", tool_registry)
        self.assertIn("retrieve_related_clauses", tool_registry)
        self.assertIn("retrieve_memory", tool_registry)
        self.assertIn("analyze_risk", tool_registry)
        self.assertIn("verify_evidence", tool_registry)
        self.assertIn("generate_revision", tool_registry)


def build_docx_bytes() -> bytes:
    document = WordDocument()
    document.add_paragraph("1. 定义")
    document.add_paragraph("保密信息是指披露方提供的商业信息。")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "披露方"
    table.cell(0, 1).text = "甲方"
    document.add_paragraph("2. 保密义务")
    document.add_paragraph("接收方应仅用于评估合作目的，并保密3年。")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_confidentiality_docx_bytes() -> bytes:
    document = WordDocument()
    document.add_paragraph("1. 保密义务")
    document.add_paragraph("接收方不得披露披露方提供的保密信息。")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_high_risk_docx_bytes() -> bytes:
    document = WordDocument()
    document.add_paragraph("保密协议")
    document.add_paragraph("披露方与接收方就保密信息签署本协议，接收方承担保密义务。")
    document.add_paragraph("1. 违约责任")
    document.add_paragraph("违约方应赔偿守约方全部损失、间接损失，且责任不限。")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_pdf_bytes() -> bytes:
    return (PROJECT_ROOT / "samples" / "pdf" / "nda_text_en.pdf").read_bytes()


if __name__ == "__main__":
    unittest.main()
